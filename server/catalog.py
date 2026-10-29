from __future__ import annotations

import json
import re
from pathlib import Path
from urllib.parse import urlparse

from server.db import connect, new_id, now, row_dict

PROFILE_COLUMNS = (
    "name_zh",
    "original_title",
    "brand",
    "category",
    "selling_points",
    "specs",
    "size",
    "package_size",
    "weight",
    "options",
    "package_contents",
    "scenes",
    "seller",
    "source_url",
    "resolved_url",
    "fetch_status",
    "missing",
)

DISPLAY = (
    ("name_zh", "中文品名"),
    ("original_title", "原文标题"),
    ("brand", "品牌 / 型号"),
    ("category", "类别"),
    ("selling_points", "核心卖点"),
    ("specs", "主要参数"),
    ("size", "产品尺寸"),
    ("package_size", "包装尺寸"),
    ("weight", "重量"),
    ("options", "规格选项"),
    ("package_contents", "包装内容"),
    ("scenes", "适用场景"),
    ("seller", "卖家"),
    ("missing", "待补"),
)

_DETAIL_COLUMNS = {
    "中文品名": "name_zh",
    "获取状态": "fetch_status",
    "原文标题": "original_title",
    "品牌 / 型号": "brand",
    "类别": "category",
    "核心卖点": "selling_points",
    "规格参数汇总": "specs",
    "产品尺寸 / 各型号尺寸": "size",
    "包装尺寸": "package_size",
    "重量（类型 / 单位）": "weight",
    "颜色 / 型号 / 规格选项": "options",
    "包装内容": "package_contents",
    "适用场景": "scenes",
    "卖家": "seller",
    "原始链接": "source_url",
    "解析后链接": "resolved_url",
    "缺失项与访问说明": "missing",
}

_SIZE_WORDS = ("尺寸", "size", "dimension", "dimensiones", "medida", "medidas")
_WEIGHT_WORDS = ("重量", "weight", "peso", "净重", "毛重")
_PACKAGE_WORDS = ("包装尺寸", "package", "empaque", "包装")
_DESC_KEYS = {"desc_detail", "description", "product_desc", "desc"}
_SELLER_KEYS = {"shop_name", "seller_name"}


_PROMPT_LABELS = ("核心卖点", "主要参数", "产品尺寸", "包装尺寸", "重量", "规格选项", "包装内容", "适用场景", "品牌 / 型号")


def save_generated_prompt(product_id: str, scene_id: str, scene_name: str, body: str, body_zh: str = "") -> None:
    body = (body or "").strip()
    scene_id = (scene_id or "").strip()
    if not product_id or not body:
        return
    scene_name = (scene_name or "").strip() or ("未接场景" if not scene_id else "未注明场景")
    with connect() as db:
        product = row_dict(db.execute("SELECT title FROM products WHERE id = ?", (product_id,)).fetchone())
        spec = row_dict(db.execute("SELECT name_zh FROM product_specs WHERE product_id = ?", (product_id,)).fetchone())
        name = ((spec or {}).get("name_zh") or (product or {}).get("title") or "未命名商品").strip()
        db.execute(
            "INSERT INTO prompts (id, name, body, body_zh, product_id, scene_id, scene_name, origin, tag, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, 'generated', '', ?)",
            (new_id("prompt"), name, body, (body_zh or "").strip(), product_id, scene_id, scene_name, now()),
        )


def page_facts(product_id: str) -> str:
    lines = []
    for block in spec_blocks(product_id):
        if block["label"] not in _PROMPT_LABELS:
            continue
        value = block["value"].strip()
        if not value or value == "页面未提供" or value.startswith("页面未提供"):
            continue
        lines.append(f"{block['label']}: {value}")
    return "\n".join(lines)


def spec_blocks(product_id: str) -> list[dict]:
    with connect() as db:
        row = row_dict(db.execute("SELECT * FROM product_specs WHERE product_id = ?", (product_id,)).fetchone())
    if not row:
        return []
    blocks = []
    for key, label in DISPLAY:
        value = (row.get(key) or "").strip()
        if value:
            blocks.append({"label": label, "value": value})
    return blocks


def save_profile(product_id: str, fields: dict, *, overwrite: bool = False) -> None:
    clean = {key: str(fields.get(key) or "").strip() for key in PROFILE_COLUMNS}
    with connect() as db:
        current = row_dict(db.execute("SELECT * FROM product_specs WHERE product_id = ?", (product_id,)).fetchone())
        if current and not overwrite:
            for key, value in clean.items():
                if not value:
                    clean[key] = current.get(key) or ""
        values = [clean[key] for key in PROFILE_COLUMNS]
        if current:
            assignments = ", ".join(f"{key} = ?" for key in PROFILE_COLUMNS)
            db.execute(
                f"UPDATE product_specs SET {assignments}, updated_at = ? WHERE product_id = ?",
                (*values, now(), product_id),
            )
        else:
            columns = ", ".join(("product_id", *PROFILE_COLUMNS, "updated_at"))
            marks = ", ".join("?" for _ in range(len(PROFILE_COLUMNS) + 2))
            db.execute(
                f"INSERT INTO product_specs ({columns}) VALUES ({marks})",
                (product_id, *values, now()),
            )
        points = clean["selling_points"]
        if points:
            db.execute(
                "UPDATE products SET points = ? WHERE id = ? AND TRIM(COALESCE(points, '')) = ''",
                (points, product_id),
            )


def _slice_between(text: str, starts: tuple[str, ...], ends: tuple[str, ...]) -> str:
    start_at = -1
    start_len = 0
    for marker in starts:
        index = text.find(marker)
        if index >= 0 and (start_at < 0 or index < start_at):
            start_at = index
            start_len = len(marker)
    if start_at < 0:
        return ""
    body = text[start_at + start_len :]
    end_at = len(body)
    for marker in ends:
        index = body.find(marker)
        if index >= 0:
            end_at = min(end_at, index)
    return re.sub(r"\n{3,}", "\n\n", body[:end_at]).strip()


def _bullet_block(text: str) -> str:
    groups: list[list[str]] = []
    current: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if re.match(r"^[✭•●▪]\s*\S", line) or re.match(r"^【[^】]{2,40}】\S", line):
            current.append(re.sub(r"^[✭•●▪]\s*", "", line).strip())
            continue
        if not line and current:
            continue
        if current:
            groups.append(current)
            current = []
    if current:
        groups.append(current)
    if not groups:
        return ""
    best = max(groups, key=len)
    if len(best) < 2:
        return ""
    return "\n".join(dict.fromkeys(item for item in best if item))


_MEASURE_UNIT = r"(?:cm|mm|in(?:ch(?:es)?)?|英寸|厘米|毫米)"
_TRIPLE = re.compile(
    rf"\d+(?:\.\d+)?\s*(?:{_MEASURE_UNIT})?\s*[x×*]\s*\d+(?:\.\d+)?\s*(?:{_MEASURE_UNIT})?\s*[x×*]\s*\d+(?:\.\d+)?\s*{_MEASURE_UNIT}",
    re.IGNORECASE,
)
_WEIGHT_VALUE = re.compile(r"\d+(?:\.\d+)?\s*(?:kg|lb|lbs|oz|公斤|千克|克|g)\b", re.IGNORECASE)
_ANY_MEASURE = re.compile(rf"\d+(?:\.\d+)?\s*(?:{_MEASURE_UNIT}|kg|lb|lbs|oz|公斤|千克|克|g)\b", re.IGNORECASE)


def _label_kind(label: str) -> str:
    lowered = label.lower()
    if _has(lowered, ("包装尺寸", "package", "empaque", "paquete", "package dimensions", "dimensiones del paquete")):
        return "package"
    if _has(lowered, _WEIGHT_WORDS) or "item weight" in lowered:
        return "weight"
    if _has(lowered, _SIZE_WORDS) or _has(lowered, ("产品尺寸", "直径", "item length", "item width", "item height", "product dimensions", "item dimensions")):
        return "size"
    return ""


def _split_measure_line(line: str) -> tuple[str, str]:
    match = re.match(r"^([^:：\n]{1,40})\s*[:：]\s*(\S.*)$", line)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    if len(line) <= 40 and _label_kind(line):
        return line, line if _ANY_MEASURE.search(line) else ""
    return "", line


def _short_measure(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    found = _TRIPLE.search(value)
    if found and len(value) > 90:
        return found.group(0)
    return value[:180]


def _measurement_fields(text: str) -> tuple[list[str], list[str], list[str]]:
    size: list[str] = []
    package: list[str] = []
    weight: list[str] = []
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    index = 0
    while index < len(lines):
        line = lines[index]
        label, value = _split_measure_line(line)
        kind = _label_kind(label)
        if kind and not _ANY_MEASURE.search(value) and index + 1 < len(lines) and _ANY_MEASURE.search(lines[index + 1]) and len(lines[index + 1]) <= 90:
            value = lines[index + 1]
            index += 1
        if not kind and _TRIPLE.search(line) and len(line) <= 90:
            kind = "size"
            value = _TRIPLE.search(line).group(0)
        if kind == "package" and _ANY_MEASURE.search(value):
            package.append(_short_measure(value))
        elif kind == "weight" and _WEIGHT_VALUE.search(value):
            weight.append(re.sub(r"\s+", " ", _WEIGHT_VALUE.search(value).group(0)))
        elif kind == "size" and _ANY_MEASURE.search(value):
            size.append(_short_measure(value))
        index += 1
    block = _slice_between(text, ("Specification:", "Specifications:", "规格参数", "产品规格"), ("Product List:", "Notes:", "注意事项", "包装清单", "Product information"))
    for raw in block.splitlines():
        line = raw.strip().lstrip(":：").strip()
        if not line or not _ANY_MEASURE.search(line):
            continue
        label, value = _split_measure_line(line)
        kind = _label_kind(label)
        if kind == "weight" and _WEIGHT_VALUE.search(value):
            weight.append(re.sub(r"\s+", " ", _WEIGHT_VALUE.search(value).group(0)))
        elif kind == "package" and _ANY_MEASURE.search(value):
            package.append(_short_measure(value))
        else:
            size.append(line[:180])
    return list(dict.fromkeys(size)), list(dict.fromkeys(package)), list(dict.fromkeys(weight))


def _attribute_lines(text: str) -> str:
    lines = []
    for raw in text.splitlines():
        line = raw.strip()
        match = re.match(r"^([^:：\n]{1,24})\s*[:：]\s*(\S.{0,180})$", line)
        if not match:
            continue
        label, value = match.group(1).strip(), re.sub(r"\s+", " ", match.group(2)).strip()
        if "$" in label or re.search(r"\d", label) or label.lower() in {"http", "https", "quantity", "qty", "数量"}:
            continue
        lines.append(f"{label}：{value}")
    return "\n".join(dict.fromkeys(lines))


def extract_profile(html: str, text: str) -> dict:
    pairs: list[tuple[str, str]] = []
    descriptions: list[str] = []
    sellers: list[str] = []
    options: list[str] = []
    for blob in _json_blobs(html):
        _walk(blob, pairs, descriptions, sellers, options)
    sizes = [value for name, value in pairs if _has(name, _SIZE_WORDS) and not _has(name, _PACKAGE_WORDS)]
    packages = [value for name, value in pairs if _has(name, _PACKAGE_WORDS)]
    weights = [value for name, value in pairs if _has(name, _WEIGHT_WORDS)]
    spec_lines = [f"{name}：{value}" for name, value in pairs if name and value]
    visible_points = _slice_between(text, ("Descripción del producto",), ("【Especificaciones Técnicas】", "Información del vendedor", "Advertencia de uso"))
    visible_specs = _slice_between(text, ("【Especificaciones Técnicas】",), ("【¿Qué incluye", "Información del vendedor"))
    visible_package = _slice_between(text, ("【¿Qué incluye el paquete?】", "¿Qué incluye el paquete?"), ("Información del vendedor", "Advertencia de uso"))
    ali_points = _slice_between(text, ("商品详情", "产品详情"), ("买家评价", "推荐货源", "看了又看"))
    selling = visible_points or _bullet_block(text) or ali_points or "\n".join(dict.fromkeys(descriptions))
    specs = visible_specs or _attribute_lines(text) or "\n".join(dict.fromkeys(spec_lines))
    text_size, text_package, text_weight = _measurement_fields(text)
    sizes = text_size or sizes
    packages = text_package or packages
    weights = text_weight or weights
    # #region agent log
    try:
        import json
        import time
        line = json.dumps({"sessionId": "02d4e0", "hypothesisId": "F", "location": "catalog.py:extract_profile", "message": "profile fields chosen", "data": {"visiblePoints": len(visible_points), "visibleSpecs": len(visible_specs), "visiblePackage": len(visible_package), "jsonDescriptions": len(descriptions)}, "timestamp": int(time.time() * 1000)}, ensure_ascii=False)
        with open(r"d:\AI\debug-02d4e0.log", "a", encoding="utf-8") as fh:
            fh.write(line + "\n")
    except Exception:
        pass
    # #endregion
    return {
        "selling_points": selling[:4000],
        "specs": specs[:8000],
        "size": "\n".join(dict.fromkeys(sizes))[:2000],
        "package_size": "\n".join(dict.fromkeys(packages))[:2000],
        "weight": "\n".join(dict.fromkeys(weights))[:1000],
        "options": "\n".join(dict.fromkeys(options))[:2000],
        "package_contents": visible_package[:2000],
        "seller": sellers[0] if sellers else "",
        "fetch_status": "已抓取" if selling or specs or options or visible_package else "",
    }


def import_workbook(path: str | Path) -> int:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    sheet = book["产品详情"] if "产品详情" in book.sheetnames else book.active
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next((index for index, row in enumerate(rows) if row and row[0] == "序号"), None)
    if header_index is None:
        raise ValueError("表格里没有产品详情")
    header = [str(cell or "").strip() for cell in rows[header_index]]
    imported = 0
    for row in rows[header_index + 1 :]:
        if not row or not row[0] or row[0] == "序号":
            continue
        fields = {column: "" for column in PROFILE_COLUMNS}
        for index, name in enumerate(header):
            key = _DETAIL_COLUMNS.get(name)
            if key and index < len(row) and row[index] is not None:
                fields[key] = str(row[index]).strip()
        if not fields["source_url"] and not fields["name_zh"]:
            continue
        product_id = _match_product(fields["source_url"], fields["resolved_url"], fields["original_title"])
        if not product_id:
            product_id = new_id("prod")
            title = fields["name_zh"] or fields["original_title"] or "未命名商品"
            with connect() as db:
                db.execute(
                    "INSERT INTO products (id, source_url, title, points, status, created_at) VALUES (?, ?, ?, ?, 'ready', ?)",
                    (product_id, fields["source_url"], title, fields["selling_points"], now()),
                )
        save_profile(product_id, fields, overwrite=True)
        imported += 1
    return imported


def _match_product(source_url: str, resolved_url: str, original_title: str) -> str:
    tokens = {token for token in (_token(source_url), _token(resolved_url)) if token}
    numbers = set(re.findall(r"\d{8,}", f"{source_url} {resolved_url}"))
    title = (original_title or "").strip().lower()
    with connect() as db:
        rows = [row_dict(row) for row in db.execute("SELECT id, source_url, title FROM products")]
    for row in rows:
        url = row["source_url"] or ""
        haystack = url.lower()
        if tokens and any(token in haystack for token in tokens):
            return row["id"]
        if numbers and any(number in url for number in numbers):
            return row["id"]
        current = (row["title"] or "").strip().lower()
        if title and len(title) > 24 and title[:40] in current:
            return row["id"]
    return ""


def _token(url: str) -> str:
    path = urlparse(url or "").path.strip("/")
    if not path:
        return ""
    return path.split("/")[-1].lower()


def _json_blobs(html: str):
    for match in re.finditer(r"<script[^>]*>(.*?)</script>", html or "", re.S | re.I):
        raw = match.group(1).strip()
        start = raw.find("{")
        if start < 0:
            continue
        try:
            yield json.loads(raw[start:])
        except json.JSONDecodeError:
            continue


def _walk(value, pairs: list[tuple[str, str]], descriptions: list[str], sellers: list[str], options: list[str]) -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            if key in _DESC_KEYS and isinstance(item, str) and item.strip():
                descriptions.append(item.strip())
            if key in _SELLER_KEYS and isinstance(item, str) and item.strip():
                sellers.append(item.strip())
        name = str(value.get("name") or value.get("prop_name") or "").strip()
        raw_value = value.get("value")
        if raw_value is None:
            raw_value = value.get("prop_value")
        if name and isinstance(raw_value, str) and raw_value.strip():
            pairs.append((name, raw_value.strip()))
        if name and isinstance(value.get("sale_prop_values"), list):
            choice = [str(item.get("prop_value")).strip() for item in value["sale_prop_values"] if isinstance(item, dict) and item.get("prop_value")]
            if choice:
                options.append(f"{name}：" + " / ".join(choice))
        for item in value.values():
            _walk(item, pairs, descriptions, sellers, options)
    elif isinstance(value, list):
        for item in value:
            _walk(item, pairs, descriptions, sellers, options)


def _has(name: str, words: tuple[str, ...]) -> bool:
    lowered = name.lower()
    return any(word.lower() in lowered for word in words)
